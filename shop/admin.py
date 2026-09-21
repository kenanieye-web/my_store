from django.contrib import admin
from .models import Category, Product, ProductImage, Cart, CartItem, Order, OrderItem, Customer
import openpyxl
from django.http import HttpResponse
from django.urls import path
from django.shortcuts import render, redirect
import requests
from django.core.files.base import ContentFile
from bs4 import BeautifulSoup
from urllib.parse import urlparse
import os


# 1. تخصيص عرض التصنيفات
@admin.register(Category)
class CategoryAdmin(admin.ModelAdmin):
    list_display = ['name', 'parent', 'slug']
    list_filter = ['parent']
    prepopulated_fields = {'slug': ('name',)}
    search_fields = ['name']


# 2. تخصيص عرض العملاء
@admin.register(Customer)
class CustomerAdmin(admin.ModelAdmin):
    list_display = ['user', 'phone', 'city', 'created_at']
    search_fields = ['user__username', 'user__first_name', 'phone', 'city']
    list_filter = ['city', 'created_at']


# 3. إتاحة رفع صور متعددة للمنتج داخل صفحة المنتج
class ProductImageInline(admin.TabularInline):
    model = ProductImage
    extra = 3
    verbose_name = "صورة إضافية"
    verbose_name_plural = "معرض الصور الإضافية"


# 4. تخصيص عرض المنتجات مع إضافة معرض الصور والبحث المتقدم واستيراد/تصدير Excel
@admin.register(Product)
class ProductAdmin(admin.ModelAdmin):
    list_display = ['name', 'category', 'price', 'stock', 'is_available', 'created_at']
    list_filter = ['is_available', 'category', 'created_at']
    list_editable = ['price', 'stock', 'is_available']
    search_fields = ['name', 'model', 'description', 'specifications']
    inlines = [ProductImageInline]
    actions = ['export_to_excel']
    change_list_template = "admin/product_changelist.html"

    EXCEL_HEADERS = [
        'ID (اتركه فارغ لمنتج جديد)', 'اسم المنتج', 'الموديل', 'التصنيف',
        'الوصف', 'المواصفات الفنية', 'السعر', 'متاح للبيع (نعم/لا)', 'المخزون',
        'رابط الصورة',
    ]

    def export_to_excel(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append(self.EXCEL_HEADERS)

        for p in queryset:
            ws.append([
                p.id,
                p.name,
                p.model or '',
                p.category.name if p.category else '',
                p.description or '',
                p.specifications or '',
                float(p.price),
                'نعم' if p.is_available else 'لا',
                p.stock,
                '',
            ])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=products.xlsx'
        wb.save(response)
        return response
    export_to_excel.short_description = "تصدير المنتجات المحددة إلى Excel"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('import-excel/', self.import_excel, name='product-import-excel'),
            path('download-template/', self.download_template, name='product-download-template'),
            path('import-from-url/', self.import_from_url, name='product-import-from-url'),
        ]
        return custom_urls + urls

    def download_template(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template"
        ws.append(self.EXCEL_HEADERS)
        ws.append(['', 'مثال: لابتوب HP', 'HP-2024', 'إلكترونيات',
                    'وصف المنتج هنا', 'المواصفات هنا', 1500, 'نعم', 10, ''])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=product_template.xlsx'
        wb.save(response)
        return response

    def import_excel(self, request):
        if request.method == "POST":
            excel_file = request.FILES["excel_file"]
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]
            errors = []

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                data = dict(zip(headers, row))
                name = data.get('اسم المنتج')
                category_name = data.get('التصنيف')

                if not name:
                    continue

                if not category_name:
                    errors.append(f"الصف {row_num}: تم تجاهله (لا يوجد تصنيف)")
                    continue

                category_obj, _ = Category.objects.get_or_create(name=category_name)

                product_obj, created = Product.objects.update_or_create(
                    id=data.get('ID (اتركه فارغ لمنتج جديد)') or None,
                    defaults={
                        'name': name,
                        'model': data.get('الموديل') or '',
                        'category': category_obj,
                        'description': data.get('الوصف') or '',
                        'specifications': data.get('المواصفات الفنية') or '',
                        'price': data.get('السعر') or 0,
                        'is_available': str(data.get('متاح للبيع (نعم/لا)')).strip() in ['نعم', 'True', 'true', '1'],
                        'stock': data.get('المخزون') or 0,
                    }
                )

                image_url = data.get('رابط الصورة')
                if image_url:
                    try:
                        img_response = requests.get(image_url, timeout=10)
                        if img_response.status_code == 200:
                            file_name = os.path.basename(urlparse(image_url).path) or f"product_{product_obj.id}.jpg"
                            product_obj.image.save(file_name, ContentFile(img_response.content), save=True)
                    except Exception as e:
                        errors.append(f"الصف {row_num}: فشل تحميل الصورة ({e})")

            msg = "تم استيراد المنتجات بنجاح"
            if errors:
                msg += " — لكن: " + " | ".join(errors)
            self.message_user(request, msg)
            return redirect("..")

        return render(request, "admin/excel_import.html")

    def import_from_url(self, request):
        if request.method == "POST":
            product_url = request.POST.get("product_url")
            if product_url:
                try:
                    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
                    response = requests.get(product_url, headers=headers, timeout=15)
                    
                    if response.status_code == 200:
                        soup = BeautifulSoup(response.text, 'html.parser')
                        
                        # سحب الاسم
                        name_tag = soup.find('meta', property='og:title')
                        product_name = name_tag['content'] if name_tag else (soup.h1.text.strip() if soup.h1 else "منتج مستورد")
                        
                        # سحب الصورة
                        image_tag = soup.find('meta', property='og:image')
                        image_url = image_tag['content'] if image_tag else ''
                        
                        # سحب السعر
                        price = 0
                        price_tag = soup.find('meta', property='product:price:amount') or soup.find(class_=['price', 'product-price', 'current-price'])
                        if price_tag:
                            price_text = price_tag['content'] if 'content' in price_tag.attrs else price_tag.text
                            import re
                            numbers = re.findall(r'\d+\.\d+|\d+', price_text.replace(',', ''))
                            if numbers:
                                price = float(numbers[0])
                        
                        default_category = Category.objects.first()
                        
                        product_obj = Product.objects.create(
                            name=product_name,
                            price=price,
                            category=default_category,
                            is_available=True,
                            stock=10,
                            description=f"تم الاستيراد من الرابط: {product_url}",
                        )
                        
                        if image_url:
                            try:
                                img_resp = requests.get(image_url, timeout=10)
                                if img_resp.status_code == 200:
                                    file_name = os.path.basename(urlparse(image_url).path) or f"product_{product_obj.id}.jpg"
                                    product_obj.image.save(file_name, ContentFile(img_resp.content), save=True)
                            except Exception:
                                pass
                        
                        self.message_user(request, f"تم سحب وإضافة المنتج ({product_name}) بنجاح من الرابط!")
                        return redirect("..")
                    else:
                        self.message_user(request, "فشل الوصول إلى الرابط المطلوب، تأكد من صحة الرابط.", level='error')
                except Exception as e:
                    self.message_user(request, f"حدث خطأ أثناء سحب البيانات: {e}", level='error')
                    
        return render(request, "admin/url_import.html")
# 5. تخصيص عرض الطلبات وعناصرها
class OrderItemInline(admin.TabularInline):
    model = OrderItem
    raw_id_fields = ['product']


@admin.register(Order)
class OrderAdmin(admin.ModelAdmin):
    list_display = ['id', 'full_name', 'city', 'phone', 'total_price', 'status', 'created_at']
    list_filter = ['status', 'created_at']
    list_editable = ['status']
    inlines = [OrderItemInline]


# 6. تسجيل باقي النماذج
admin.site.register(Cart)
admin.site.register(CartItem)