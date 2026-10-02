from django.contrib import admin
from decimal import Decimal
from django.db.models import DecimalField, ExpressionWrapper, F, Value
from shop.models import Category, PricingSettings, Product, ProductImage
import openpyxl
from django.http import HttpResponse
from django.urls import path
from django.shortcuts import render, redirect
import requests
from django.core.files.base import ContentFile
from bs4 import BeautifulSoup
from urllib.parse import urlparse
import os
import zipfile

from shop.admin.product_image_inline import ProductImageInline


@admin.register(Product)
class ProductAdmin(admin.ModelAdmin):
    list_display = [
        'name', 'category', 'cost_price',
        'retail_price_col', 'wholesale_price_col', 'super_wholesale_price_col',
        'stock', 'is_available', 'created_at',
    ]
    list_filter = ['is_available', 'category', 'created_at']
    list_editable = ['cost_price', 'stock', 'is_available']
    search_fields = ['name', 'model', 'description', 'specifications']
    inlines = [ProductImageInline]
    actions = ['export_to_excel']
    change_list_template = "admin/product_changelist.html"

    # ===== أعمدة الأسعار المحسوبة (تُحسب في قاعدة البيانات باستعلام واحد) =====
    def get_queryset(self, request):
        qs = super().get_queryset(request)
        pricing = PricingSettings.get_solo()

        def with_margin(margin):
            factor = Value(
                Decimal('1') + Decimal(margin) / Decimal('100'),
                output_field=DecimalField(max_digits=10, decimal_places=4),
            )
            return ExpressionWrapper(
                F('cost_price') * factor,
                output_field=DecimalField(max_digits=14, decimal_places=2),
            )

        return qs.annotate(
            retail_calc=with_margin(pricing.retail_margin_percent),
            wholesale_calc=with_margin(pricing.wholesale_margin_percent),
            super_wholesale_calc=with_margin(pricing.super_wholesale_margin_percent),
        )

    @staticmethod
    def _fmt(value):
        return '-' if value is None else f"{value:.2f}"

    @admin.display(description='سعر التجزئة', ordering='retail_calc')
    def retail_price_col(self, obj):
        return self._fmt(getattr(obj, 'retail_calc', None))

    @admin.display(description='سعر الجملة', ordering='wholesale_calc')
    def wholesale_price_col(self, obj):
        return self._fmt(getattr(obj, 'wholesale_calc', None))

    @admin.display(description='سعر جملة الجملة', ordering='super_wholesale_calc')
    def super_wholesale_price_col(self, obj):
        return self._fmt(getattr(obj, 'super_wholesale_calc', None))

    EXCEL_HEADERS = [
        'ID (اتركه فارغ لمنتج جديد)', 'اسم المنتج', 'الموديل', 'التصنيف',
        'الوصف', 'المواصفات الفنية', 'سعر التكلفة', 'متاح للبيع (نعم/لا)', 'المخزون',
        'رابط الصورة (اختياري)', 'صورة المنتج', 'صورة العلبة',
    ]

    def export_to_excel(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append(self.EXCEL_HEADERS)

        for p in queryset:
            ws.append([
                p.id,
                p.name,
                p.model or '',
                p.category.name if p.category else '',
                p.description or '',
                p.specifications or '',
                float(p.cost_price),
                'نعم' if p.is_available else 'لا',
                p.stock,
                '',
                '',  # صورة المنتج
                '',  # صورة العلبة
            ])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=products.xlsx'
        wb.save(response)
        return response
    export_to_excel.short_description = "تصدير المنتجات المحددة إلى Excel"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('import-excel/', self.import_excel, name='import-excel'),
            path('import-excel/', self.import_excel, name='shop_product_import-excel'),
            
            path('download-template/', self.download_template, name='download-template'),
            path('download-template/', self.download_template, name='shop_product_download-template'),
            
            path('import-from-url/', self.import_from_url, name='import-from-url'),
            path('import-from-url/', self.import_from_url, name='shop_product_import-from-url'),
        ]
        return custom_urls + urls

    def download_template(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template"
        ws.append(self.EXCEL_HEADERS)
        ws.append(['', 'مثال: لابتوب HP', 'HP-2024', 'إلكترونيات',
                    'وصف المنتج هنا', 'المواصفات هنا', 1000, 'نعم', 10, '', '', ''])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=product_template.xlsx'
        wb.save(response)
        return response

    def import_excel(self, request):
        if request.method == "POST":
            excel_file = request.FILES.get("excel_file")
            if not excel_file:
                self.message_user(request, "الرجاء رفع ملف إكسل صالح.", level='error')
                return redirect("..")

            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
            
            zip_file = request.FILES.get("zip_file")
            images_dict = {}
            if zip_file:
                with zipfile.ZipFile(zip_file, 'r') as z:
                    for filename in z.namelist():
                        if not filename.endswith('/') and '__MACOSX' not in filename:
                            images_dict[os.path.basename(filename)] = z.read(filename)
                            
            headers = [cell.value for cell in ws[1]]
            errors = []

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                # تخطي الصفوف الفارغة تماماً التي لا تحتوي على أي بيانات
                if not any(row):
                    continue

                data = dict(zip(headers, row))
                name = data.get('اسم المنتج')
                category_name = data.get('التصنيف')

                # تنظيف النصوص من الفراغات الزائدة إن وجدت
                if name:
                    name = str(name).strip()
                if category_name:
                    category_name = str(category_name).strip()

                if not name:
                    errors.append(f"الصف {row_num}: تم تجاهله لعدم وجود اسم للمنتج.")
                    continue

                if not category_name:
                    errors.append(f"الصف {row_num}: تم تجاهله لعدم وجود تصنيف للمنتج ({name}).")
                    continue

                category_obj, _ = Category.objects.get_or_create(name=category_name)

                # معالجة الآيدي إن وجد وفارغ
                raw_id = data.get('ID (اتركه فارغ لمنتج جديد)')
                product_id = int(raw_id) if raw_id and str(raw_id).isdigit() else None

                product_obj, created = Product.objects.update_or_create(
                    id=product_id,
                    defaults={
                        'name': name,
                        'model': str(data.get('الموديل') or '').strip(),
                        'category': category_obj,
                        'description': str(data.get('الوصف') or '').strip(),
                        'specifications': str(data.get('المواصفات الفنية') or '').strip(),
                        'cost_price': data.get('سعر التكلفة') or data.get('السعر') or 0,  # «السعر» لتوافق الملفات القديمة
                        'is_available': str(data.get('متاح للبيع (نعم/لا)')).strip() in ['نعم', 'True', 'true', '1', 'Yes', 'yes'],
                        'stock': data.get('المخزون') or 0,
                    }
                )

                product_img_name = data.get('صورة المنتج')
                package_img_name = data.get('صورة العلبة')

                def find_image_content(img_name):
                    if not img_name:
                        return None, None
                    name_str = str(img_name).strip()
                    if name_str in images_dict:
                        return name_str, images_dict[name_str]
                    for ext in ['.jpg', '.JPG', '.png', '.PNG', '.jpeg', '.JPEG']:
                        if (name_str + ext) in images_dict:
                            return name_str + ext, images_dict[name_str + ext]
                    return None, None

                found_name, img_bytes = find_image_content(product_img_name)
                if found_name and img_bytes:
                    product_obj.image.save(
                        found_name,
                        ContentFile(img_bytes),
                        save=True
                    )

                found_pkg_name, pkg_bytes = find_image_content(package_img_name)
                if found_pkg_name and pkg_bytes:
                    gallery_image = ProductImage(product=product_obj)
                    gallery_image.image.save(
                        found_pkg_name,
                        ContentFile(pkg_bytes),
                        save=True
                    )

                image_url = data.get('رابط الصورة (اختياري)') or data.get('رابط الصورة')
                if image_url:
                    try:
                        img_response = requests.get(str(image_url).strip(), timeout=10)
                        if img_response.status_code == 200:
                            file_name = os.path.basename(urlparse(str(image_url).strip()).path) or f"product_{product_obj.id}.jpg"
                            product_obj.image.save(file_name, ContentFile(img_response.content), save=True)
                    except Exception as e:
                        errors.append(f"الصف {row_num}: فشل تحميل الصورة من الرابط للمنتج {name} ({e})")

            msg = "تم استيراد المنتجات بنجاح."
            if errors:
                msg += " ملاحظات وتنبيهات: " + " | ".join(errors)
                self.message_user(request, msg, level='warning')
            else:
                self.message_user(request, msg, level='success')
                
            return redirect("..")

        return render(request, "admin/excel_import.html")

    def import_from_url(self, request):
        if request.method == "POST":
            product_url = request.POST.get("product_url")
            if product_url:
                try:
                    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
                    response = requests.get(product_url, headers=headers, timeout=15)
                    
                    if response.status_code == 200:
                        soup = BeautifulSoup(response.text, 'html.parser')
                        
                        name_tag = soup.find('meta', property='og:title')
                        product_name = name_tag['content'] if name_tag else (soup.h1.text.strip() if soup.h1 else "منتج مستورد")
                        
                        image_tag = soup.find('meta', property='og:image')
                        image_url = image_tag['content'] if image_tag else ''
                        
                        price = 0
                        price_tag = soup.find('meta', property='product:price:amount') or soup.find(class_=['price', 'product-price', 'current-price'])
                        if price_tag:
                            price_text = price_tag['content'] if 'content' in price_tag.attrs else price_tag.text
                            import re
                            numbers = re.findall(r'\d+\.\d+|\d+', price_text.replace(',', ''))
                            if numbers:
                                price = float(numbers[0])
                        
                        default_category = Category.objects.first()
                        
                        product_obj = Product.objects.create(
                            name=product_name,
                            cost_price=price,
                            category=default_category,
                            is_available=True,
                            stock=10,
                            description=f"تم الاستيراد من الرابط: {product_url}",
                        )
                        
                        if image_url:
                            try:
                                img_resp = requests.get(image_url, timeout=10)
                                if img_resp.status_code == 200:
                                    file_name = os.path.basename(urlparse(image_url).path) or f"product_{product_obj.id}.jpg"
                                    product_obj.image.save(file_name, ContentFile(img_resp.content), save=True)
                            except Exception:
                                pass
                        
                        self.message_user(request, f"تم سحب وإضافة المنتج ({product_name}) بنجاح من الرابط!")
                        return redirect("..")
                    else:
                        self.message_user(request, "فشل الوصول إلى الرابط المطلوب، تأكد من صحة الرابط.", level='error')
                except Exception as e:
                    self.message_user(request, f"حدث خطأ أثناء سحب البيانات: {e}", level='error')
                    
        return render(request, "admin/url_import.html")